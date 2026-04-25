from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
from bs4 import BeautifulSoup
from docx import Document

from .ocr import OcrUnavailableError, ocr_pdf


@dataclass
class ExtractedDocument:
    title: str
    text: str
    source_type: str


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".doc",
}


def extract_document(file_path: str | Path, title: str | None = None) -> ExtractedDocument:
    path = Path(file_path).resolve()
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")

    derived_title = title or path.stem.replace("_", " ").replace("-", " ").strip() or path.stem

    if suffix in {".txt", ".md", ".markdown"}:
        text = path.read_text(encoding="utf-8")
        return ExtractedDocument(title=derived_title, text=text, source_type="note")

    if suffix in {".html", ".htm"}:
        html = path.read_text(encoding="utf-8")
        text = BeautifulSoup(html, "html.parser").get_text("\n")
        return ExtractedDocument(title=derived_title, text=text, source_type="web")

    if suffix == ".pdf":
        return ExtractedDocument(title=derived_title, text=_extract_pdf(path), source_type="pdf")

    if suffix == ".docx":
        return ExtractedDocument(title=derived_title, text=_extract_docx(path), source_type="docx")

    if suffix == ".doc":
        return ExtractedDocument(title=derived_title, text=_extract_doc_via_libreoffice(path), source_type="doc")

    raise ValueError(f"Unhandled file type: {suffix}")


def copy_original(file_path: str | Path, target_dir: str | Path) -> Path:
    path = Path(file_path).resolve()
    target = Path(target_dir) / path.name
    shutil.copy2(path, target)
    return target


def _extract_pdf(path: Path) -> str:
    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            text = page.get_text("text")
            if text:
                parts.append(text.strip())
    text = "\n\n".join(part for part in parts if part).strip()
    if text:
        return text
    try:
        return ocr_pdf(path)
    except OcrUnavailableError as exc:
        raise ValueError(f"No embedded PDF text found and OCR fallback is unavailable for {path}: {exc}") from exc


def _extract_docx(path: Path) -> str:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    if not text.strip():
        raise ValueError(f"No text extracted from DOCX: {path}")
    return text


def _extract_doc_via_libreoffice(path: Path) -> str:
    with tempfile.TemporaryDirectory() as tmpdir:
        outdir = Path(tmpdir)
        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to",
            "txt:Text",
            str(path),
            "--outdir",
            str(outdir),
        ]
        proc = subprocess.run(cmd, check=False, capture_output=True, text=True)
        if proc.returncode != 0:
            raise ValueError(f"LibreOffice conversion failed for {path}: {proc.stderr.strip() or proc.stdout.strip()}")
        txt_path = outdir / f"{path.stem}.txt"
        if not txt_path.exists():
            raise ValueError(f"LibreOffice did not produce text output for {path}")
        text = txt_path.read_text(encoding="utf-8", errors="ignore")
        if not text.strip():
            raise ValueError(f"No text extracted from DOC: {path}")
        return text
