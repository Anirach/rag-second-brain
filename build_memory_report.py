#!/usr/bin/env python3
"""
Arthur's Knowledge & Memory System — Complete Reference
Professional DOCX report generator
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import copy

# ── Color palette ──────────────────────────────────────────────────
NAVY      = RGBColor(0x1B, 0x3A, 0x5C)   # H1
BLUE      = RGBColor(0x2A, 0x64, 0x96)   # H2
DARKGRAY  = RGBColor(0x2C, 0x3E, 0x50)   # H3
BODY_COLOR= RGBColor(0x33, 0x33, 0x33)   # body text
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_NAVY= RGBColor(0x1B, 0x3A, 0x5C)   # table header bg
TABLE_ALT = RGBColor(0xF2, 0xF6, 0xFA)   # alternating row
ACCENT    = RGBColor(0x2A, 0x64, 0x96)   # accent line

OUTPUT_PATH = "/home/clawdbot/clawd/tmp/Arthur_Memory_System_v1.docx"

# ── Helpers ────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    """Set cell background shading."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    """Set cell border."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), kwargs.get(edge, 'none'))
        tag.set(qn('w:sz'), kwargs.get('sz', '4'))
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), kwargs.get('color', '1B3A5C'))
        tcBorders.append(tag)
    tcPr.append(tcBorders)

def add_page_number_footer(doc):
    """Add 'Page X of Y' footer to all sections."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.clear()

        # Separator line above footer
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        top = OxmlElement('w:top')
        top.set(qn('w:val'), 'single')
        top.set(qn('w:sz'), '6')
        top.set(qn('w:space'), '1')
        top.set(qn('w:color'), '2A6496')
        pBdr.append(top)
        pPr.append(pBdr)

        run = para.add_run("Page ")
        run.font.name = 'Arial'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        run._r.append(fldChar1)

        instrText = OxmlElement('w:instrText')
        instrText.text = 'PAGE'
        run._r.append(instrText)

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar2)

        run2 = para.add_run(" of ")
        run2.font.name = 'Arial'
        run2.font.size = Pt(9)
        run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'begin')
        run2._r.append(fldChar3)

        instrText2 = OxmlElement('w:instrText')
        instrText2.text = 'NUMPAGES'
        run2._r.append(instrText2)

        fldChar4 = OxmlElement('w:fldChar')
        fldChar4.set(qn('w:fldCharType'), 'end')
        run2._r.append(fldChar4)


def add_page_break(doc):
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(docx_break_type())
    return para

def docx_break_type():
    from docx.enum.text import WD_BREAK
    return WD_BREAK.PAGE


def apply_body_style(para, size=11, italic=False, bold=False, color=None):
    """Apply standard body text formatting to a paragraph."""
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '80')
    spacing.set(qn('w:after'), '80')
    spacing.set(qn('w:line'), '276')
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)
    for run in para.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.color.rgb = color or BODY_COLOR
        run.font.italic = italic
        run.font.bold = bold


def add_body(doc, text, italic=False, bold=False, color=None, size=11):
    """Add a justified body paragraph."""
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.color.rgb = color or BODY_COLOR
    run.font.italic = italic
    run.font.bold = bold
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = para.paragraph_format
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(4)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.15
    return para


def add_h1(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = NAVY
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = para.paragraph_format
    fmt.space_before = Pt(18)
    fmt.space_after = Pt(8)
    fmt.keep_with_next = True
    # Bottom border
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1B3A5C')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return para


def add_h2(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = BLUE
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = para.paragraph_format
    fmt.space_before = Pt(14)
    fmt.space_after = Pt(5)
    fmt.keep_with_next = True
    return para


def add_h3(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = DARKGRAY
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = para.paragraph_format
    fmt.space_before = Pt(10)
    fmt.space_after = Pt(4)
    fmt.keep_with_next = True
    return para


def add_bullet(doc, text, level=0, bold_prefix=None):
    """Add a bullet point. Optional bold_prefix for 'Term: explanation' style."""
    para = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run1 = para.add_run(bold_prefix)
        run1.font.name = 'Arial'
        run1.font.size = Pt(11)
        run1.font.bold = True
        run1.font.color.rgb = BODY_COLOR
        run2 = para.add_run(text)
        run2.font.name = 'Arial'
        run2.font.size = Pt(11)
        run2.font.color.rgb = BODY_COLOR
    else:
        run = para.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.color.rgb = BODY_COLOR
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = para.paragraph_format
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(2)
    if level > 0:
        fmt.left_indent = Inches(0.25 * (level + 1))
    return para


def add_code_block(doc, code_text):
    """Add a monospace code block."""
    para = doc.add_paragraph()
    run = para.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = para.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.left_indent = Inches(0.3)
    # Light background
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EEF4FA')
    pPr.append(shd)
    return para


def add_table(doc, headers, rows, col_widths=None):
    """Add a professional styled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header row
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bg(cell, '1B3A5C')
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(h)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = WHITE

    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        bg = 'F2F6FA' if ri % 2 == 0 else 'FFFFFF'
        for ci, cell_text in enumerate(row_data):
            cell = row.cells[ci]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_bg(cell, bg)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(str(cell_text))
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = BODY_COLOR

    # Column widths
    if col_widths:
        for row in table.rows:
            for ci, width in enumerate(col_widths):
                row.cells[ci].width = Inches(width)

    return table


def add_info_box(doc, title, content_lines, color_hex='2A6496'):
    """Add a shaded info box with a title."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, 'EEF4FA')
    # Border
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '8')
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), color_hex)
        tcBorders.append(tag)
    tcPr.append(tcBorders)

    cell.width = Inches(6.0)
    p = cell.paragraphs[0]
    run = p.add_run(title)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = BLUE
    for line in content_lines:
        np = cell.add_paragraph()
        run2 = np.add_run(line)
        run2.font.name = 'Arial'
        run2.font.size = Pt(10)
        run2.font.color.rgb = BODY_COLOR
    return table


# ── Cover Page ─────────────────────────────────────────────────────

def build_cover(doc):
    # Top accent bar via table
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, '1B3A5C')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(' ')
    run.font.size = Pt(14)
    table.rows[0].height = Cm(1.2)

    doc.add_paragraph()  # spacer

    # Title block
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("Arthur's Knowledge &")
    run.font.name = 'Arial'
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY
    title_para.paragraph_format.space_after = Pt(2)

    title_para2 = doc.add_paragraph()
    title_para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = title_para2.add_run("Memory System")
    run2.font.name = 'Arial'
    run2.font.size = Pt(26)
    run2.font.bold = True
    run2.font.color.rgb = NAVY
    title_para2.paragraph_format.space_after = Pt(8)

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run("Complete Reference")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(16)
    sub_run.font.bold = False
    sub_run.font.color.rgb = BLUE
    sub_para.paragraph_format.space_after = Pt(30)

    # Decorative divider
    div_para = doc.add_paragraph()
    div_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div_run = div_para.add_run('— — — — — — — — — — — — —')
    div_run.font.name = 'Arial'
    div_run.font.size = Pt(12)
    div_run.font.color.rgb = BLUE
    div_para.paragraph_format.space_after = Pt(30)

    # Metadata table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    meta_data = [
        ('Prepared for', 'Anirach Mongkolboriboon'),
        ('System', 'Arthur (OpenClaw AI Assistant)'),
        ('Version', 'v1.0'),
        ('Date', datetime.now().strftime('%B %d, %Y')),
        ('Classification', 'Internal Reference'),
    ]

    for i, (label, value) in enumerate(meta_data):
        row = meta_table.rows[i]
        # Label cell
        lc = row.cells[0]
        set_cell_bg(lc, 'EEF4FA')
        lp = lc.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lr = lp.add_run(label + ':')
        lr.font.name = 'Arial'
        lr.font.size = Pt(11)
        lr.font.bold = True
        lr.font.color.rgb = NAVY
        lc.width = Inches(2.0)

        # Value cell
        vc = row.cells[1]
        set_cell_bg(vc, 'FFFFFF')
        vp = vc.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        vr = vp.add_run(value)
        vr.font.name = 'Arial'
        vr.font.size = Pt(11)
        vr.font.color.rgb = BODY_COLOR
        vc.width = Inches(3.5)

    doc.add_paragraph()

    # Bottom accent bar
    bot_table = doc.add_table(rows=1, cols=1)
    bot_cell = bot_table.rows[0].cells[0]
    set_cell_bg(bot_cell, '2A6496')
    bp = bot_cell.paragraphs[0]
    bp.paragraph_format.space_before = Pt(0)
    bp.paragraph_format.space_after = Pt(0)
    bot_run = bp.add_run(' ')
    bot_run.font.size = Pt(8)
    bot_table.rows[0].height = Cm(0.5)

    # Page break after cover
    doc.add_page_break()


# ── Main builder ───────────────────────────────────────────────────

def build_report():
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.page_width  = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin  = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin   = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    # Default paragraph font
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    # ── Cover ──────────────────────────────────────────────────────
    build_cover(doc)

    # ── Table of Contents placeholder ─────────────────────────────
    add_h1(doc, "Table of Contents")
    toc_items = [
        ("1.", "Executive Summary", "3"),
        ("2.", "System Architecture Overview", "3"),
        ("3.", "Core Memory Files", "4"),
        ("4.", "Daily Memory System", "6"),
        ("5.", "Obsidian Vault (Long-Term Memory)", "7"),
        ("6.", "SQLite Memory Database", "8"),
        ("7.", "Knowledge Graph System", "9"),
        ("8.", "Search & Retrieval", "10"),
        ("9.", "Session Learning Pipeline (/remember)", "11"),
        ("10.", "Auto-Reflection System", "12"),
        ("11.", "Proactive Memory (Heartbeat System)", "13"),
        ("12.", "Cron Jobs — Memory-Related Automation", "14"),
        ("13.", "Security & Privacy", "15"),
        ("14.", "Agent Teams & Memory", "16"),
        ("15.", "Arscontexta Inspiration", "17"),
        ("16.", "Quick Reference — Commands", "18"),
    ]
    toc_table = doc.add_table(rows=len(toc_items), cols=3)
    toc_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (num, title, pg) in enumerate(toc_items):
        row = toc_table.rows[i]
        bg = 'F2F6FA' if i % 2 == 0 else 'FFFFFF'
        # num
        set_cell_bg(row.cells[0], bg)
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(num)
        r0.font.name = 'Arial'; r0.font.size = Pt(11); r0.font.bold = True; r0.font.color.rgb = NAVY
        row.cells[0].width = Inches(0.4)
        # title
        set_cell_bg(row.cells[1], bg)
        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(title)
        r1.font.name = 'Arial'; r1.font.size = Pt(11); r1.font.color.rgb = BODY_COLOR
        row.cells[1].width = Inches(5.1)
        # page
        set_cell_bg(row.cells[2], bg)
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(pg)
        r2.font.name = 'Arial'; r2.font.size = Pt(11); r2.font.color.rgb = BLUE
        row.cells[2].width = Inches(0.5)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 1 — EXECUTIVE SUMMARY
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "1. Executive Summary")
    add_body(doc,
        "Arthur is an AI assistant powered by OpenClaw, deployed for Anirach Mongkolboriboon "
        "(university lecturer and AI engineer, Bangkok, Thailand). Unlike a typical chatbot that "
        "forgets everything between sessions, Arthur is equipped with a sophisticated multi-layer "
        "knowledge and memory architecture designed to deliver persistent, context-aware assistance "
        "across days, weeks, and months.")
    add_body(doc,
        "The system integrates three memory tiers — ephemeral short-term context, structured daily "
        "logs, and a rich long-term knowledge base — with automated pipelines that continuously "
        "distil insights from sessions into durable records. Key components include:")
    add_bullet(doc, "Core identity and context files (SOUL.md, USER.md, MEMORY.md, AGENTS.md)")
    add_bullet(doc, "A daily memory scratchpad (memory/YYYY-MM-DD.md) auto-synced to the Obsidian vault")
    add_bullet(doc, "An Obsidian-based long-term knowledge graph with wiki-linked notes and MOCs")
    add_bullet(doc, "A SQLite database (memory.db) holding 210+ indexed memories with sub-4ms FTS")
    add_bullet(doc, "Session learning and auto-reflection pipelines running on cron schedules")
    add_bullet(doc, "A proactive heartbeat system for email, calendar, and agent monitoring")
    add_bullet(doc, "Five specialised agent teams with shared Obsidian workspaces")
    add_body(doc,
        "This document provides a complete technical reference for the entire system, covering "
        "architecture, file formats, tooling, automation schedules, security boundaries, and "
        "operational command reference.")

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # SECTION 2 — SYSTEM ARCHITECTURE
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "2. System Architecture Overview")

    add_h2(doc, "2.1 Three-Layer Memory Model")
    add_body(doc,
        "Arthur's memory is structured as three distinct but interconnected tiers, each serving a "
        "different time horizon and access pattern:")

    layers = [
        ("Layer 1 — Short-Term (In-Session)",
         "Short-term",
         "Active LLM context window",
         "Current conversation only",
         "Loaded at session start from core files; lost when session ends"),
        ("Layer 2 — Daily (Operational)",
         "Daily",
         "memory/YYYY-MM-DD.md + heartbeat-state.json + running-log.md",
         "Days to weeks",
         "Raw session logs, scratch notes; auto-synced to Obsidian Daily/"),
        ("Layer 3 — Long-Term (Permanent)",
         "Long-term",
         "Obsidian Vault + SQLite memory.db + MEMORY.md",
         "Weeks to years",
         "Curated knowledge, decisions, lessons; indexed for fast retrieval"),
    ]

    headers = ["Layer", "Store", "Retention", "Description"]
    rows = [
        (l[1], l[2], l[3], l[4]) for l in layers
    ]
    add_table(doc, headers, rows, col_widths=[1.0, 2.0, 1.2, 2.3])
    doc.add_paragraph()

    add_h2(doc, "2.2 File System Layout")
    add_body(doc,
        "The workspace root is /home/clawdbot/clawd/. Key directories and files are organised as follows:")
    add_code_block(doc,
        "/home/clawdbot/clawd/\n"
        "├── SOUL.md               # Arthur's identity & persona\n"
        "├── USER.md               # Anirach's profile, preferences, timezone\n"
        "├── MEMORY.md             # Long-term curated memory (main session only)\n"
        "├── AGENTS.md             # Workspace rules, agent team definitions\n"
        "├── HEARTBEAT.md          # Proactive check rules\n"
        "├── IDENTITY.md           # Name, creature, vibe\n"
        "├── memory/\n"
        "│   ├── YYYY-MM-DD.md     # Daily session logs\n"
        "│   ├── heartbeat-state.json  # Last check timestamps\n"
        "│   └── running-log.md    # Auto-updated activity record\n"
        "├── memory.db             # SQLite indexed memories\n"
        "├── tools/\n"
        "│   ├── memory_db.py      # SQLite memory CLI\n"
        "│   ├── quick_note.py     # Route notes to correct store\n"
        "│   ├── session_remember.py  # Session learning miner\n"
        "│   ├── auto_reflection.py   # Nightly reflection\n"
        "│   ├── kg_builder.py     # Knowledge graph builder\n"
        "│   ├── kg_query.py       # Knowledge graph query\n"
        "│   └── obsidian_search.py   # Full vault search\n"
        "└── gdrive/               # Google Drive upload helpers\n"
        "\n"
        "/home/clawdbot/obsidian-vault/\n"
        "├── KnowledgeGraph/       # People, Projects, Topics, Decisions…\n"
        "├── Agents/               # Team workspaces (Writing/Academic/etc.)\n"
        "├── Daily/                # Synced daily notes\n"
        "├── Quick-Reference.md    # Critical IDs, links, facts\n"
        "└── knowledge_graph.json  # Graph data file")

    add_h2(doc, "2.3 How Layers Connect and Sync")
    add_body(doc,
        "The three layers are kept in sync through a combination of automated cron jobs, "
        "session-triggered writes, and heartbeat callbacks:")
    add_bullet(doc, "During a session: Arthur writes observations directly to today's memory/YYYY-MM-DD.md and uses quick_note.py to route structured notes to the vault.")
    add_bullet(doc, "At 23:00 BKK: session_remember.py scans today's file, extracts structured learnings via DeepSeek LLM, and writes to MEMORY.md, KnowledgeGraph/Decisions/, and Action-Items/.")
    add_bullet(doc, "At 02:00 BKK: auto_reflection.py reviews the last two days, distils lessons, updates running-log.md and MEMORY.md, and commits the vault to GitHub.")
    add_bullet(doc, "3× daily: kg_builder.py rebuilds the knowledge graph from vault notes.")
    add_bullet(doc, "4× daily: the vault is git-pushed to GitHub for off-site backup.")
    add_bullet(doc, "Weekly (Sunday 03:00): memory-maintenance script prunes and reorganises MEMORY.md.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 3 — CORE MEMORY FILES
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "3. Core Memory Files")
    add_body(doc,
        "Six core files are loaded at the start of each session (main session loads all; "
        "sub-agents and group-chat sessions load a restricted subset for security). "
        "Together they provide Arthur's identity, user context, and operational rules.")

    add_h2(doc, "3.1 MEMORY.md — Long-Term Curated Memory")
    add_body(doc,
        "MEMORY.md is Arthur's primary long-term memory store — the distilled essence of "
        "everything learned across all past sessions. It is a curated markdown document updated "
        "by both Arthur and the automated pipelines.")
    add_h3(doc, "Loading Policy")
    add_body(doc,
        "MEMORY.md is loaded ONLY in the main session (Anirach's direct Telegram chat). "
        "It is never loaded in group chats, Discord servers, or sessions involving other "
        "participants. This policy protects sensitive personal context from leaking to "
        "third parties.")
    add_h3(doc, "Content Categories")
    add_bullet(doc, "Lessons Learned — insights extracted from past conversations and mistakes.")
    add_bullet(doc, "Decisions — key architectural and strategic choices with rationale.")
    add_bullet(doc, "Preferences — Anirach's communication style, tool preferences, language choices.")
    add_bullet(doc, "Active Projects — brief status of ongoing work.")
    add_bullet(doc, "Facts — important identifiers, credentials, and reference data.")
    add_h3(doc, "Update Flow")
    add_body(doc,
        "Arthur can freely read and edit MEMORY.md during main sessions. The session_remember.py "
        "cron appends new learnings nightly. The memory-maintenance cron reviews and prunes "
        "outdated entries weekly. Unlike daily files (raw logs), MEMORY.md is always curated "
        "content — the 'wisdom' not the 'diary'.")

    doc.add_paragraph()
    add_h2(doc, "3.2 SOUL.md — Identity and Persona")
    add_body(doc,
        "SOUL.md defines who Arthur is: personality traits, communication style, values, "
        "working preferences, and the 'spirit' guiding interactions. It is loaded at the "
        "very start of every session as the foundational identity layer.")
    add_bullet(doc, "Personality: curious, warm, direct, occasionally witty — like a trusted colleague.")
    add_bullet(doc, "Communication: concise by default, detailed when asked; never padded.")
    add_bullet(doc, "Work ethic: completes tasks without constant check-ins; asks only when genuinely uncertain.")
    add_bullet(doc, "Values: accuracy, privacy, human oversight, continuous learning.")
    add_body(doc,
        "Arthur is permitted and encouraged to edit SOUL.md over time as personality and "
        "style evolve through experience — making this a living document rather than a static config.")

    doc.add_paragraph()
    add_h2(doc, "3.3 USER.md — Anirach's Profile")
    add_body(doc,
        "USER.md contains a comprehensive profile of Anirach Mongkolboriboon — Arthur's "
        "primary user. This file is the foundation for personalised, context-aware assistance.")
    add_h3(doc, "Key Data Points")
    cols = ["Category", "Details"]
    rows_u = [
        ("Full Name", "Anirach Mongkolboriboon"),
        ("Role", "University Lecturer & AI Engineer"),
        ("Location", "Bangkok, Thailand (Asia/Bangkok, UTC+7)"),
        ("Language", "Thai (primary) + English (technical)"),
        ("Research Interests", "AI in healthcare, RAG systems, LLM applications, higher education"),
        ("Communication Style", "Direct; prefers bullet points over walls of text"),
        ("Working Hours", "Flexible; often works late evenings Bangkok time"),
        ("AI Tools", "OpenClaw, DeepSeek, Claude, Obsidian"),
    ]
    add_table(doc, cols, rows_u, col_widths=[2.0, 4.5])

    doc.add_paragraph()
    add_h2(doc, "3.4 AGENTS.md — Workspace Rules & Team Definitions")
    add_body(doc,
        "AGENTS.md is the operational rulebook for Arthur and all spawned sub-agents. "
        "It defines workflow patterns, memory conventions, safety rules, and agent team structures. "
        "Every sub-agent receives AGENTS.md as its primary operating context.")
    add_h3(doc, "Key Sections in AGENTS.md")
    add_bullet(doc, "First Run instructions (bootstrap flow)")
    add_bullet(doc, "Session startup checklist (what to read, in what order)")
    add_bullet(doc, "Memory system usage guide (when to use vault vs MEMORY.md vs daily files)")
    add_bullet(doc, "Agent team definitions and spawn triggers for Coding, Writing, Academic, Translation, Course teams")
    add_bullet(doc, "Research parallel-spawn strategy")
    add_bullet(doc, "Safety rules (de-identification, file protection, destructive command safeguards)")
    add_bullet(doc, "Google Drive file naming convention")
    add_bullet(doc, "Heartbeat guidelines and check intervals")

    doc.add_paragraph()
    add_h2(doc, "3.5 HEARTBEAT.md — Proactive Check Rules")
    add_body(doc,
        "HEARTBEAT.md is a lightweight checklist read by Arthur during each heartbeat poll "
        "(triggered by OpenClaw approximately every 30 minutes). It specifies what to check, "
        "under what conditions to alert, and what background work can be done silently.")
    add_bullet(doc, "Email inbox — alert if urgent unread messages detected.")
    add_bullet(doc, "Calendar — alert if events are within 2 hours.")
    add_bullet(doc, "Weather — flag if relevant to planned outdoor activity.")
    add_bullet(doc, "Active sub-agents — check for stalled or completed jobs.")
    add_body(doc,
        "The cardinal rule: respond HEARTBEAT_OK when nothing actionable exists. "
        "Avoid noise. Check state tracking in heartbeat-state.json to prevent "
        "repeat notifications for the same event.")

    doc.add_paragraph()
    add_h2(doc, "3.6 IDENTITY.md — Name, Creature, Vibe")
    add_body(doc,
        "IDENTITY.md holds Arthur's public-facing identity specifics: the chosen name (Arthur), "
        "the associated creature (a golden Labrador 🐕), and the overall personality vibe. "
        "This file anchors the consistent persona across all channels and contexts — ensuring "
        "Arthur presents the same 'face' whether in Telegram, Discord, or spawned as a sub-agent.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 4 — DAILY MEMORY SYSTEM
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "4. Daily Memory System")
    add_body(doc,
        "The daily memory layer provides a raw operational journal — a rolling record of "
        "what happened each day across all sessions. It acts as the source material for the "
        "automated learning pipelines that distil insights into long-term memory.")

    add_h2(doc, "4.1 Daily Session Logs")
    add_body(doc,
        "Each day gets its own markdown file at:")
    add_code_block(doc, "/home/clawdbot/clawd/memory/YYYY-MM-DD.md")
    add_body(doc,
        "Arthur writes to this file throughout the day — logging key events, decisions made, "
        "tasks completed, and observations worth remembering. The format is intentionally loose "
        "(markdown bullet points, timestamps, freeform notes) to avoid friction during sessions. "
        "These files are auto-synced to the Obsidian vault at Daily/YYYY-MM-DD.md by cron.")
    add_h3(doc, "Typical Content")
    add_bullet(doc, "Session timestamps and attendees")
    add_bullet(doc, "Tasks completed and tools used")
    add_bullet(doc, "Decisions made (with brief rationale)")
    add_bullet(doc, "Problems encountered and solutions found")
    add_bullet(doc, "Links, file paths, and identifiers referenced")
    add_bullet(doc, "Ideas and observations worth capturing")

    add_h2(doc, "4.2 heartbeat-state.json — Check Time Tracker")
    add_body(doc,
        "heartbeat-state.json is a small JSON file that records the Unix timestamp of the "
        "last time each proactive check was performed. This prevents Arthur from re-alerting "
        "on the same email or calendar event across consecutive heartbeats.")
    add_code_block(doc,
        "{\n"
        '  "lastChecks": {\n'
        '    "email": 1703275200,\n'
        '    "calendar": 1703260800,\n'
        '    "weather": null,\n'
        '    "sub_agents": 1703270000\n'
        "  }\n"
        "}")
    add_body(doc,
        "The heartbeat system reads this file before each check to determine whether enough "
        "time has elapsed since the last check of that type. Checks are skipped if performed "
        "within the last 30 minutes.")

    add_h2(doc, "4.3 running-log.md — Activity Record")
    add_body(doc,
        "running-log.md is an auto-updated chronological activity record — a higher-level "
        "summary than the raw daily files. It is updated by the auto_reflection.py cron and "
        "serves as a quick-glance history of what Arthur has been doing over recent weeks. "
        "Useful for onboarding new sessions quickly without reading every daily file.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 5 — OBSIDIAN VAULT
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "5. Obsidian Vault (Long-Term Memory)")
    add_body(doc,
        "The Obsidian vault is Arthur's primary long-term knowledge base — a structured, "
        "wiki-linked collection of notes covering people, projects, decisions, research topics, "
        "and agent workspaces. It is THE source of truth for persistent knowledge.")
    add_code_block(doc, "Location: /home/clawdbot/obsidian-vault/")

    add_h2(doc, "5.1 Vault Structure")
    add_code_block(doc,
        "/home/clawdbot/obsidian-vault/\n"
        "├── KnowledgeGraph/\n"
        "│   ├── People/           # Contacts, collaborators\n"
        "│   ├── Projects/         # Project notes and status\n"
        "│   ├── Topics/           # Research topics\n"
        "│   ├── Decisions/        # Architecture & strategy decisions\n"
        "│   ├── Action-Items/     # Tasks (Pending.md, etc.)\n"
        "│   ├── Documents/        # Facts, references\n"
        "│   ├── Events/           # Milestones, submissions, meetings\n"
        "│   └── MOC-*.md          # Maps of Content (hub, domain, topic)\n"
        "├── Agents/\n"
        "│   ├── Shared/           # Cross-team resources & templates\n"
        "│   ├── Writing/          # Writing Team workspace\n"
        "│   ├── Academic/         # Academic Team workspace\n"
        "│   ├── Translation/      # Translation Team workspace\n"
        "│   ├── Course/           # Course Team workspace\n"
        "│   └── Coding/           # Coding Team workspace\n"
        "├── Daily/\n"
        "│   └── YYYY-MM-DD.md     # Synced from memory/ files\n"
        "├── Quick-Reference.md    # Critical IDs, links, facts\n"
        "└── knowledge_graph.json  # Computed graph data")

    add_h2(doc, "5.2 Wiki-Links and Connectivity")
    add_body(doc,
        "All notes use Obsidian-style wiki-links ([[Note Name]]) to create a connected "
        "knowledge graph. Links are automatically updated by kg_auto_link.py when new notes "
        "are added. This enables graph-based exploration of related concepts — e.g., clicking "
        "from a Project note to related People, Decisions, and Action-Items.")
    add_bullet(doc, "[[Person]] links connect projects to collaborators")
    add_bullet(doc, "[[Decision]] links show why architectural choices were made")
    add_bullet(doc, "[[Action-Items/Pending]] aggregates all outstanding tasks")
    add_bullet(doc, "[[MOC-Research]] provides a navigable hub for all research notes")

    add_h2(doc, "5.3 Standard Tags")
    headers_t = ["Tag", "Meaning", "When to Apply"]
    rows_t = [
        ("#project", "Active project", "When creating or updating a project note"),
        ("#complete", "Completed work", "When a project or task is finished"),
        ("#decision", "Architecture/design decision", "When recording a significant choice"),
        ("#research", "Research notes", "All academic and exploratory content"),
        ("#todo", "Pending tasks", "Action items that need follow-up"),
        ("#blocked", "Blocked items", "Tasks waiting on external dependencies"),
    ]
    add_table(doc, headers_t, rows_t, col_widths=[1.2, 2.0, 3.3])

    add_h2(doc, "5.4 Maps of Content (MOCs)")
    add_body(doc,
        "Maps of Content are hub notes that aggregate links to related notes at three levels: "
        "Hub MOCs (entire domains), Domain MOCs (sub-topics), and Topic MOCs (specific areas). "
        "Examples include MOC-Research.md, MOC-Infrastructure.md, and MOC-Teaching.md. "
        "MOCs are the navigational backbone of the vault.")

    add_h2(doc, "5.5 Git Synchronisation")
    add_body(doc,
        "The vault is a git repository pushed to GitHub. Automated git commits and pushes "
        "occur 4× daily (02:00, 08:00, 14:00, 20:00 BKK) via cron. Arthur also commits "
        "immediately after significant updates (new decisions, project milestones, etc.). "
        "This provides off-site backup and version history.")

    add_h2(doc, "5.6 Quick-Reference.md")
    add_body(doc,
        "Quick-Reference.md is the vault's 'cheat sheet' — a single file loaded at every "
        "session start containing critical identifiers, frequently used links, API keys "
        "locations, important contacts, and facts that need instant access. "
        "It is the first file Arthur reads after core identity files.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 6 — SQLITE MEMORY DATABASE
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "6. SQLite Memory Database")
    add_body(doc,
        "The SQLite memory database provides fast, structured, full-text searchable storage "
        "for discrete memory units — facts, decisions, lessons, and observations. Unlike the "
        "Obsidian vault (optimised for linked documents) or MEMORY.md (curated narrative), "
        "memory.db is optimised for machine-readable queries and rapid retrieval.")
    add_code_block(doc,
        "Location: /home/clawdbot/clawd/memory.db\n"
        "CLI:      python3 tools/memory_db.py\n"
        "Size:     210+ indexed memories")

    add_h2(doc, "6.1 Key Features")
    add_bullet(doc, "Full-text search (FTS5) across all memory content")
    add_bullet(doc, "Sub-4ms query latency on typical hardware")
    add_bullet(doc, "Structured fields: id, type, content, tags, created_at, updated_at, archived")
    add_bullet(doc, "Archiving support — outdated memories archived rather than deleted")
    add_bullet(doc, "Export to JSON for backup or analysis")

    add_h2(doc, "6.2 CLI Commands")
    headers_db = ["Command", "Example", "Description"]
    rows_db = [
        ("add", "memory_db.py add --type fact 'text'", "Add a new memory"),
        ("search", "memory_db.py search 'RAG system'", "Full-text search"),
        ("recent", "memory_db.py recent --days 7", "List recent memories"),
        ("list", "memory_db.py list --type decision", "List by type"),
        ("get", "memory_db.py get 42", "Get memory by ID"),
        ("update", "memory_db.py update 42 'new text'", "Update memory content"),
        ("archive", "memory_db.py archive 42", "Archive a memory"),
        ("stats", "memory_db.py stats", "Show database statistics"),
        ("export", "memory_db.py export out.json", "Export all memories to JSON"),
        ("rebuild-fts", "memory_db.py rebuild-fts", "Rebuild full-text search index"),
    ]
    add_table(doc, headers_db, rows_db, col_widths=[1.3, 2.7, 2.5])

    add_h2(doc, "6.3 Memory Types")
    add_bullet(doc, "fact — Specific identifiable facts (IDs, URLs, specs, credentials)")
    add_bullet(doc, "decision — Architecture or strategy decisions with rationale")
    add_bullet(doc, "lesson — Things learned from experience or mistakes")
    add_bullet(doc, "todo — Actionable tasks to track")
    add_bullet(doc, "observation — General observations and notes")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 7 — KNOWLEDGE GRAPH SYSTEM
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "7. Knowledge Graph System")
    add_body(doc,
        "The Knowledge Graph system transforms the flat collection of Obsidian notes into a "
        "queryable network graph — enabling relationship-based queries like 'What projects is "
        "person X involved in?' or 'Which decisions affect this topic?'")

    add_h2(doc, "7.1 Graph Tools")
    headers_kg = ["Tool", "Location", "Purpose"]
    rows_kg = [
        ("kg_builder.py", "tools/kg_builder.py", "Builds the graph from vault notes by parsing wiki-links and frontmatter"),
        ("kg_query.py", "tools/kg_query.py", "Query the graph: find related notes, shortest paths, node neighbours"),
        ("kg_auto_link.py", "tools/kg_auto_link.py", "Scans new notes and auto-inserts wiki-links to existing related notes"),
        ("knowledge_graph.json", "obsidian-vault/knowledge_graph.json", "The serialised graph data (nodes + edges + metadata)"),
    ]
    add_table(doc, headers_kg, rows_kg, col_widths=[1.8, 2.0, 2.7])

    add_h2(doc, "7.2 Graph Data Format")
    add_body(doc,
        "knowledge_graph.json stores the graph as a JSON object with nodes (each note) and "
        "edges (wiki-link relationships). Each node carries metadata: title, tags, creation "
        "date, and note type. Edges are directed (source → target) with link type annotation.")
    add_code_block(doc,
        '{\n'
        '  "nodes": [\n'
        '    {"id": "Projects/RAG-System", "title": "RAG System", "tags": ["#project"], ...}\n'
        '  ],\n'
        '  "edges": [\n'
        '    {"source": "Projects/RAG-System", "target": "People/Anirach", "type": "involves"}\n'
        '  ]\n'
        '}')

    add_h2(doc, "7.3 Rebuild Schedule")
    add_body(doc,
        "The graph is rebuilt 3× daily by cron (01:00, 06:00, 11:00 BKK) via a cron job "
        "running kg_builder.py. Arthur can also trigger a manual rebuild at any time by "
        "running the script directly. kg_auto_link.py runs after each vault write to keep "
        "links fresh.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 8 — SEARCH & RETRIEVAL
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "8. Search & Retrieval")
    add_body(doc,
        "Arthur has three distinct search methods, each optimised for different scenarios. "
        "Choosing the right method ensures the best trade-off between recall, speed, and precision.")

    add_h2(doc, "8.1 Method Comparison")
    headers_s = ["Method", "Tool", "Best For", "Latency"]
    rows_s = [
        ("Semantic Search", "memory_search", "Finding related concepts even when wording differs; cross-session context", "~200ms"),
        ("SQLite FTS", "memory_db.py search", "Exact term search across structured memories; fast lookups", "<4ms"),
        ("Vault Search", "obsidian_search.py", "Full-text search across all vault notes; finding documents by content", "~50ms"),
    ]
    add_table(doc, headers_s, rows_s, col_widths=[1.5, 1.8, 2.8, 0.9])

    add_h2(doc, "8.2 memory_search — Semantic Search")
    add_body(doc,
        "memory_search is OpenClaw's built-in semantic search tool. It uses vector embeddings "
        "to find relevant content even when the exact words differ. It indexes MEMORY.md and "
        "all files in memory/. Best used when you need to find something but aren't sure of "
        "the exact phrasing — e.g., 'What did we decide about the course structure?'")
    add_bullet(doc, "Scope: MEMORY.md + memory/YYYY-MM-DD.md files")
    add_bullet(doc, "Returns: ranked snippets with file path + line numbers")
    add_bullet(doc, "When to use: vague queries, cross-session context, preference retrieval")

    add_h2(doc, "8.3 memory_db.py search — SQLite FTS")
    add_body(doc,
        "SQLite full-text search provides sub-4ms exact-term matching across the 210+ "
        "structured memories in memory.db. Use this when you know what you're looking for "
        "and want the fastest possible response.")
    add_code_block(doc,
        "python3 tools/memory_db.py search \"RAG system\"\n"
        "python3 tools/memory_db.py search \"hospital accreditation\" --type decision")
    add_bullet(doc, "Scope: all records in memory.db")
    add_bullet(doc, "Returns: matching rows with type, content, date")
    add_bullet(doc, "When to use: exact keyword lookup, type-filtered searches, recent memory queries")

    add_h2(doc, "8.4 obsidian_search.py — Full Vault Search")
    add_body(doc,
        "obsidian_search.py scans all markdown files in the Obsidian vault for text matches. "
        "It is the broadest search — covering research notes, project files, people profiles, "
        "and agent workspaces that may not be in memory.db.")
    add_code_block(doc,
        "python3 tools/obsidian_search.py \"Arscontexta\"\n"
        "python3 tools/obsidian_search.py \"#project\" --tag")
    add_bullet(doc, "Scope: all .md files in /home/clawdbot/obsidian-vault/")
    add_bullet(doc, "Returns: matching file paths + content snippets")
    add_bullet(doc, "When to use: finding documents, research notes, vault-wide queries")

    add_h2(doc, "8.5 Recommended Search Strategy")
    add_body(doc, "Follow this decision flow for optimal retrieval:")
    add_bullet(doc, "1. Start with memory_search for broad, natural-language queries about past sessions.")
    add_bullet(doc, "2. Use memory_db.py search for fast, specific term lookups on structured memories.")
    add_bullet(doc, "3. Fall back to obsidian_search.py for vault-wide document search.")
    add_bullet(doc, "4. Use memory_get (OpenClaw) to read specific files after locating them via search.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 9 — SESSION LEARNING PIPELINE
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "9. Session Learning Pipeline (/remember)")
    add_body(doc,
        "The session learning pipeline is a core innovation inspired by the Arscontexta "
        "methodology. It automatically mines each day's session logs for structured "
        "learnings and distributes them to the appropriate long-term stores.")
    add_code_block(doc, "Script: tools/session_remember.py")

    add_h2(doc, "9.1 Three Operating Modes")
    headers_m = ["Mode", "Command", "Description"]
    rows_m = [
        ("Direct text", "session_remember.py \"text to remember\"", "Process a specific piece of text immediately"),
        ("Stdin", "session_remember.py --stdin", "Read input from stdin (pipe from other commands)"),
        ("Auto", "session_remember.py --auto", "Scan today's memory/YYYY-MM-DD.md and extract all learnings"),
    ]
    add_table(doc, headers_m, rows_m, col_widths=[1.2, 2.8, 2.5])

    add_h2(doc, "9.2 Extraction Categories")
    add_body(doc,
        "The pipeline extracts five categories of structured learning from raw session text:")
    add_bullet(doc, "Lessons — experiential insights ('I learned that X causes Y')")
    add_bullet(doc, "Decisions — explicit choices with context ('We decided to use DeepSeek because...')")
    add_bullet(doc, "TODOs — actionable tasks identified during the session")
    add_bullet(doc, "Facts — specific identifiable facts (IDs, names, URLs, specs)")
    add_bullet(doc, "Connections — relationships between concepts or entities")

    add_h2(doc, "9.3 LLM Extraction with Fallback")
    add_body(doc,
        "Primary extraction uses DeepSeek LLM via API — it is instructed to output structured "
        "JSON with category-labelled extractions. If the API is unavailable or returns an error, "
        "the script falls back to a heuristic regex parser that catches common patterns like "
        "'TODO:', 'Decision:', 'Note:', and 'Learned:' prefix markers.")

    add_h2(doc, "9.4 Output Destinations")
    add_bullet(doc, "Daily log: appends a '## Session Learnings' section to today's memory file")
    add_bullet(doc, "MEMORY.md: appends new lessons and decisions under appropriate sections")
    add_bullet(doc, "KnowledgeGraph/Decisions/: creates individual note files for each decision")
    add_bullet(doc, "KnowledgeGraph/Action-Items/Pending.md: appends new TODOs")

    add_h2(doc, "9.5 Cron Schedule")
    add_body(doc,
        "session_remember.py --auto runs nightly at 23:00 Bangkok time via cron, ensuring "
        "every day's learnings are captured before midnight. Arthur can also trigger it "
        "manually mid-session with specific text.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # SECTION 10 — AUTO-REFLECTION SYSTEM
    # ══════════════════════════════════════════════════════════════
    add_h1(doc, "10. Auto-Reflection System")
    add_body(doc,
        "The auto-reflection system runs a deeper nightly review, looking across multiple "
        "days of memory files to identify patterns, extract lessons, and maintain the "
        "running activity log. It is the 'overnight thinking' layer.")
    add_code_block(doc,
        "Script:   tools/auto_reflection.py\n"
        "Schedule: Daily at 02:00 Bangkok (Asia/Bangkok)")

    add_h2(doc, "10.1 Process Flow")
    add_bullet(doc, "1. Read last 2 days of memory/YYYY-MM-DD.md files")
    add_bullet(doc, "2. Send combined content to DeepSeek LLM with reflection prompt")
    add_bullet(doc, "3. LLM extracts: key events, lessons learned, unresolved questions, progress")
    add_bullet(doc, "4. Fallback to heuristic extraction if LLM unavailable")
    add_bullet(doc, "5. Append distilled summary to running-log.md")
    add_bullet(doc, "6. Update MEMORY.md with new lessons (no duplicates)")
    add_bullet(doc, "7. Git commit and push vault to GitHub")

    add_h2(doc, "10.2 Running-Log.md Format")
    add_body(doc,
        "running-log.md uses a dated entry format with structured sections per reflection:")
    add_code_block(doc,
        "## 2026-02-23 (Auto-Reflection)\n"
        "### Key Events\n"
        "- Built Arthur_Memory_System_v1.docx report\n"
        "- Deployed session_remember.py cron job\n"
        "### Lessons Learned\n"
        "- DeepSeek API faster than expected for extraction tasks\n"
        "### Open Questions\n"
        "- Should kg_auto_link.py run more frequently?")

    add_h2(doc, "10.3 MEMORY.md Deduplication")
    add_body(doc,
        "Before appending to MEMORY.md, auto_reflection.py checks for semantic similarity "
        "with existing entries to avoid duplicating lessons already recorded. "
        "It uses simple string matching (not embedding comparison) — a future enhancement "
        "could use the ChromaDB vector store for smarter deduplication.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
