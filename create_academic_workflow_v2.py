#!/usr/bin/env python3
"""Create Academic Paper Production Workflow v2 DOCX report using the formatting standards."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy

# Colors matching the template
NAVY = RGBColor(0x1B, 0x3A, 0x5C)
BLUE = RGBColor(0x2A, 0x64, 0x96)
DARK_GRAY = RGBColor(0x2C, 0x3E, 0x50)
GRAY = RGBColor(0x66, 0x66, 0x66)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HEADER_BG = "1B3A5C"
TABLE_ALT_BG = "F8F9FA"
BLACK = RGBColor(0x33, 0x33, 0x33)

doc = Document()

# Set margins (1 inch = 914400 EMU)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_font(run, name='Arial', size=None, bold=None, color=None):
    run.font.name = name
    if size: run.font.size = size
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = color

def add_cover_line(text, size, bold=None, color=NAVY, spacing_after=Pt(0)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = spacing_after
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p

def add_heading1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=Pt(16), bold=True, color=NAVY)
    # Add bottom border
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="1B3A5C"/></w:pBdr>')
    pPr.append(pBdr)
    return p

def add_heading2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=Pt(13), bold=True, color=BLUE)
    return p

def add_heading3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=Pt(12), bold=True, color=DARK_GRAY)
    return p

def add_body(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=Pt(11), bold=True, color=BLACK)
        r2 = p.add_run(text)
        set_font(r2, size=Pt(11), color=BLACK)
    else:
        r = p.add_run(text)
        set_font(r, size=Pt(11), color=BLACK)
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    if bold_prefix:
        r0 = p.add_run("• ")
        set_font(r0, size=Pt(11), color=BLACK)
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=Pt(11), bold=True, color=BLACK)
        r2 = p.add_run(text)
        set_font(r2, size=Pt(11), color=BLACK)
    else:
        r = p.add_run("• " + text)
        set_font(r, size=Pt(11), color=BLACK)
    return p

def add_numbered(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(f"{num}. {text}")
    set_font(r, size=Pt(11), color=BLACK)
    return p

def set_cell_shading(cell, color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_font(r, size=Pt(10), bold=True, color=WHITE)
        set_cell_shading(cell, TABLE_HEADER_BG)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Data rows
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_font(r, size=Pt(10), color=BLACK)
            if ri % 2 == 1:
                set_cell_shading(cell, TABLE_ALT_BG)
    
    # Add borders
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)
    
    doc.add_paragraph()  # spacing
    return table

# ==================== COVER PAGE ====================
doc.add_paragraph()  # spacer
doc.add_paragraph()  # spacer
doc.add_paragraph()  # spacer

add_cover_line("Academic Paper Production Workflow v2", Pt(24), bold=True, color=NAVY, spacing_after=Pt(8))
add_cover_line("From Idea to Submission in 7 Days", Pt(18), color=BLUE, spacing_after=Pt(8))
add_cover_line("Lessons Learned from 23 Versions of MKG-RAG Second Brain Paper", Pt(15), color=GRAY, spacing_after=Pt(24))

# Divider line
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("━" * 40)
set_font(r, size=Pt(12), color=ORANGE)

add_cover_line("Prepared by: Arthur AI Agent System", Pt(11), color=GRAY, spacing_after=Pt(4))
add_cover_line("Date: February 17, 2026", Pt(11), color=GRAY, spacing_after=Pt(4))
add_cover_line("Version 2.0  |  Internal Guide", Pt(10), color=GRAY)

# Page break
doc.add_page_break()

# ==================== 1. EXECUTIVE SUMMARY ====================
add_heading1("1. Executive Summary")
add_body("This workflow replaces the ad-hoc process that produced 23 versions over 3 weeks. Based on lessons from the MKG-RAG Second Brain paper journey, it targets 3 drafts in 7 days with higher novelty and fewer iterations.")
add_body("Key improvements: experiments-first approach, upfront novelty validation, venue-first planning, batched reviews, and quality gates from day 1.")

# ==================== 2. PROBLEMS WITH THE OLD WORKFLOW ====================
add_heading1("2. Problems with the Old Workflow")

add_heading2("2.1 Experiments Came Too Late")
add_body("v1-v9 built on proxy/synthetic claims. Real experiments (DDXPlus) didn't arrive until v15-v17. Papers without real results get destroyed in review.")

add_heading2("2.2 No Novelty Validation Upfront")
add_body('"Second Brain" was marketing until v23 when grounded in Extended Mind theory. Reviewers see through ungrounded framing immediately.')

add_heading2("2.3 Quality System Was Reactive")
add_body("QUALITY_LOCKS and 3-agent review gate created AFTER the v9 crisis (score dropped from 8.5 to 4.4). Should have been there from the start.")

add_heading2("2.4 Too Many Small Iterations")
add_body("23 versions, most were incremental patches. Each review-fix cycle costs time and introduces regression risk.")

add_heading2("2.5 Venue Selected Late")
add_body("Started targeting Information Fusion, pivoted to AIiH. Format conversion (IEEE to LNCS) wasted multiple versions.")

add_heading2("2.6 Theoretical Grounding Added Last")
add_body("Clark & Chalmers Extended Mind, Hutchins Distributed Cognition added in final versions. Should anchor novelty claims from the beginning.")

# ==================== 3. THE NEW WORKFLOW ====================
add_heading1("3. The New Workflow: 3 Phases, 7 Days")

add_heading2("3.1 Phase 1: Foundation (Day 1-2)")
add_body("Tasks:")
add_bullet("Pick target venue (format, page limit, scope, deadline)")
add_bullet("Literature gap analysis: define novel contribution in ONE sentence")
add_bullet('Run "Does this exist?" check (Semantic Scholar, Google Scholar, arXiv)')
add_bullet("Build baseline experiments with real data")
add_bullet("Design ablation-ready architecture (each component removable)")
add_bullet("Create comparison table: what existing methods do vs what yours adds")
add_bullet("Anchor method in existing theory (not just a name — cite cognitive science, information theory, etc.)")

add_body("Deliverables table:")
add_table(
    ["Deliverable", "Description", "Owner"],
    [
        ["Venue Brief", "Target venue, format, page limit, scope", "Journal Scout"],
        ["Novelty Statement", "One sentence, grounded in theory", "Paper Architect"],
        ["Literature Gap Table", "Existing methods vs yours", "Literature Lead"],
        ["Baseline Results", "Real experiments, real data", "Data Analyst"],
        ["Ablation Design", "Component removal plan", "Methodology Expert"],
    ]
)

add_heading2("3.2 Phase 2: Full Draft (Day 3-5)")
add_body("Tasks:")
add_bullet("Write complete paper with real results (not placeholders)")
add_bullet("All sections: intro, related work, method, experiments, discussion, conclusion")
add_bullet("Include all tables, figures, equations")
add_bullet("Apply venue format from the start (e.g., llncs.cls, IEEE template)")
add_bullet("3-agent parallel review (peer-reviewer, methodology-expert, technical-writer)")
add_bullet("Batch ALL fixes into single revision (Draft 2)")

add_body("Review gate table:")
add_table(
    ["Reviewer", "Focus", "Pass Criteria"],
    [
        ["Peer Reviewer", "Baselines, novelty, claims vs evidence", "No unsupported claims"],
        ["Methodology Expert", "Statistical rigor, reproducibility", "Proper CIs, tests, caveats"],
        ["Technical Writer", "Clarity, flow, readability", "Score ≥ 4.0/5.0"],
    ]
)

add_body('Rule: ALL THREE must pass. No exceptions. No "ship it with one approval."')

add_heading2("3.3 Phase 3: Polish & Submit (Day 6-7)")
add_body("Tasks:")
add_bullet("Apply all Draft 2 review fixes")
add_bullet("Final 3-agent review (should be minor issues only)")
add_bullet("Format compliance check (page count, references, anonymization)")
add_bullet("Generate supplementary materials if needed")
add_bullet("Submit")

# ==================== 4. QUALITY GATES ====================
add_heading1("4. Quality Gates (Active from Day 1)")

add_heading3("4.1 QUALITY_LOCKS")
add_body("Lock any element praised by reviewers. Never regress.")

add_heading3("4.2 POSITIVE_FEEDBACK tracker")
add_body("Cumulative praise log across review rounds.")

add_heading3("4.3 Three-Agent Review Gate")
add_body("Every draft reviewed by 3 agents in parallel. Unanimous approval required.")

add_heading3("4.4 Novelty Checkpoint")
add_body("Before Phase 2, verify the one-sentence novelty claim still holds after literature review.")

add_heading3("4.5 Claims-Evidence Alignment")
add_body("Every claim in the paper must map to a specific experiment, table, or citation. No orphan claims.")

# ==================== 5. NOVELTY BOOSTERS ====================
add_heading1("5. Novelty Boosters")

add_heading3("5.1 Ablation-First Design")
add_body('Build experiments so each component can be independently removed. Creates "unique contribution per component" — reviewers always ask for this.')

add_heading3("5.2 Theoretical Grounding Early")
add_body("Don't just name your method — anchor it in established theory from draft 1. Examples: Extended Mind (Clark & Chalmers), Distributed Cognition (Hutchins), Information Foraging Theory.")

add_heading3("5.3 Comparison Table in Outline")
add_body("Before writing, create a table of existing methods and what yours uniquely adds. If no clear gap, rethink the contribution.")

add_heading3("5.4 Real-World Framing")
add_body('Position contribution as solving a real problem, not just improving metrics. "Why does this matter to a doctor/engineer/user?"')

add_heading3("5.5 Honest Limitations")
add_body("Proactively state what your method does NOT do. Builds credibility and prevents reviewer attacks.")

# ==================== 6. AGENT TEAM ROLES ====================
add_heading1("6. Agent Team Roles")

add_body("Table of all 9 agents:")
add_table(
    ["Agent", "Phase", "Role", "Model"],
    [
        ["Paper Architect", "All", "Principal investigator, strategy", "Opus"],
        ["Literature Lead", "Phase 1", "Gap analysis, related work", "Opus"],
        ["Journal Scout", "Phase 1", "Venue selection, format requirements", "Opus"],
        ["Data Analyst", "Phase 1-2", "Experiments, statistics, ablations", "Opus"],
        ["Methodology Expert", "Phase 2-3", "Statistical rigor review", "Opus"],
        ["Technical Writer", "Phase 2-3", "Writing quality, clarity", "Sonnet"],
        ["Peer Reviewer", "Phase 2-3", "Tough conference reviewer simulation", "Opus"],
        ["Ethics Reviewer", "Phase 2", "Dataset ethics, bias, double-blind", "Opus"],
        ["Format Editor", "Phase 3", "Venue format compliance", "Sonnet"],
    ]
)

# ==================== 7. CHECKLIST ====================
add_heading1("7. Checklist: Before Submitting")

add_body("Numbered checklist:")
add_numbered(1, "Novelty claim is one sentence, grounded in theory")
add_numbered(2, "All experiments use real data (no proxies)")
add_numbered(3, "Ablation study shows each component's unique contribution")
add_numbered(4, "Statistical tests applied (Wilson CIs, McNemar's, etc.)")
add_numbered(5, "External baselines included (not just your own variants)")
add_numbered(6, "All 3 reviewers gave ACCEPT")
add_numbered(7, "Page count within venue limit")
add_numbered(8, "Double-blind compliance (no author names, no self-citations revealing identity)")
add_numbered(9, "References complete and correctly formatted")
add_numbered(10, "Supplementary materials prepared (code repo, appendix)")

# ==================== 8. TIMELINE COMPARISON ====================
add_heading1("8. Timeline Comparison")

add_body("Table:")
add_table(
    ["Aspect", "Old (MKG-RAG)", "New (v2 Workflow)"],
    [
        ["Duration", "~21 days", "7 days target"],
        ["Versions", "23", "3 drafts"],
        ["Real experiments", "Added at v15", "From Day 1"],
        ["Novelty validation", "v23 (last minute)", "Day 1"],
        ["Venue format", "Converted mid-way", "Applied from start"],
        ["Quality gates", "Added at v9 crisis", "Active from Day 1"],
        ["Review rounds", "10+", "2 (Phase 2 + Phase 3)"],
    ]
)

# ==================== 9. LESSONS ====================
add_heading1("9. Lessons from the MKG-RAG Journey")

add_body("Bullet points:")
add_bullet("Honesty > Impressive Claims: ", "Conceptual papers can be publishable without fake results")
add_bullet("Quality Locks Prevent Regression: ", "23 locks ensured improvements were cumulative")
add_bullet("Three Parallel Reviewers catch different issues: ", "baselines, statistics, readability")
add_bullet("ML Baselines Beating RAG Is Okay: ", "Frame as accuracy vs generalization trade-off")
add_bullet("Medical Datasets > General Benchmarks: ", "Domain-specific data makes contribution tangible")
add_bullet("Verify Numbers Before Locking: ", "Wrong numbers hurt more than vague ones")
add_bullet("Venue Selection Matters: ", "Focused venue (AIiH) better than broad (Information Fusion)")
add_bullet('"Second Brain" Grounding: ', "Marketing language became theoretical contribution when grounded in Extended Mind thesis")
add_bullet("Never deliver without full team review: ", "One-agent papers get rejected")

# Add page numbers in footer
for section in doc.sections:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r = p.add_run("Page ")
    set_font(r, size=Pt(9), color=GRAY)
    
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    r1 = p.add_run()
    r1._r.append(fldChar1)
    
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    r2 = p.add_run()
    r2._r.append(instrText)
    
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r3 = p.add_run()
    r3._r.append(fldChar2)
    
    r4 = p.add_run(" of ")
    set_font(r4, size=Pt(9), color=GRAY)
    
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    r5 = p.add_run()
    r5._r.append(fldChar3)
    
    instrText2 = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> NUMPAGES </w:instrText>')
    r6 = p.add_run()
    r6._r.append(instrText2)
    
    fldChar4 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r7 = p.add_run()
    r7._r.append(fldChar4)

outpath = "Academic_Paper_Production_Workflow_v2.docx"
doc.save(outpath)
print(f"Saved to {outpath}")