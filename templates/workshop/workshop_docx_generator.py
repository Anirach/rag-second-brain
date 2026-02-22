#!/usr/bin/env python3
"""
Workshop DOCX Generator
========================
Generates professional workshop documents matching the Vibe Coding style.

Design System:
- Cover: Purple title (40pt), gray subtitle/description, purple accent bar
- Part headers: Purple background (#7C3AED) with white text
- Module headers: H1 (18pt bold), with colored duration badge
- Section headers: H2 (14pt bold)
- Body: 12pt, justified
- AI Prompt boxes: Green header (#10B981), dark code block (#1E1E1E), light tip (#ECFDF5)
- Best Practice boxes: Green background (#D4EDDA)
- Vibe/Tip boxes: Purple background (#FDF4FF)
- Warning boxes: Yellow background (#FFF3CD)
- Tables: Purple header row, alternating shading
- Bullets: List Paragraph style

Usage:
    python3 workshop_docx_generator.py config.json output.docx
    
Config JSON structure:
{
    "title": "🎵 VIBE CODING",
    "subtitle": "with Claude Code Extension",
    "description": "Complete Team Development Workshop",
    "tech_stack": ["React 18 • TypeScript 5", "Tailwind CSS • Claude Code Extension"],
    "duration": "8+ Hours",
    "team_size": "3-6 Developers",
    "features": "📋 Step-by-Step Guides • AI Prompts • Team Workflows",
    "banner": "🤖 Copy-Paste AI Prompts for Every Task",
    "parts": [
        {
            "title": "PART I: FOUNDATION",
            "subtitle": "Understanding Fundamentals",
            "modules": [
                {
                    "title": "Module 1: Introduction",
                    "duration": "45 minutes",
                    "format": "Concept Introduction",
                    "badge_color": "purple",
                    "sections": [
                        {
                            "heading": "1.1 What is Vibe Coding?",
                            "content": [
                                {"type": "text", "value": "Body text here..."},
                                {"type": "vibe", "value": "🎵 VIBE: Core Philosophy..."},
                                {"type": "best_practice", "value": "✅ BEST PRACTICE: ..."},
                                {"type": "warning", "value": "⚠️ WARNING: ..."},
                                {"type": "bullets", "items": ["Item 1", "Item 2"]},
                                {"type": "code", "value": "const x = 1;"},
                                {"type": "ai_prompt", "title": "Your First Prompt",
                                 "prompt": "@file src/... - Create...",
                                 "tip": "AI will read the file and..."},
                                {"type": "table", "headers": ["Col1", "Col2"],
                                 "rows": [["A", "B"], ["C", "D"]]},
                                {"type": "numbered_table", "headers": ["#", "Step"],
                                 "rows": [["1", "Do this"], ["2", "Do that"]]},
                                {"type": "quick_ref", "title": "Prompt Title",
                                 "prompt": "The actual prompt text..."}
                            ]
                        }
                    ]
                }
            ]
        }
    ],
    "summary_cards": [
        {"title": "🎵 VIBE CODING", "items": ["Point 1", "Point 2"]},
        {"title": "🎨 DESIGN SYSTEM", "items": ["Point 1", "Point 2"]}
    ]
}
"""

import json
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ─── Color Palette ───
PURPLE       = RGBColor(0x7C, 0x3A, 0xED)
PURPLE_DARK  = RGBColor(0x5B, 0x21, 0xB6)
GREEN        = RGBColor(0x10, 0xB9, 0x81)
GREEN_DARK   = RGBColor(0x06, 0x5F, 0x46)
GREEN_TEXT   = RGBColor(0x15, 0x57, 0x24)
GRAY_900     = RGBColor(0x1F, 0x29, 0x37)
GRAY_700     = RGBColor(0x37, 0x41, 0x51)
GRAY_500     = RGBColor(0x6B, 0x72, 0x80)
GRAY_400     = RGBColor(0x9C, 0xA3, 0xAF)
GRAY_200     = RGBColor(0xE5, 0xE7, 0xEB)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
BLACK        = RGBColor(0x1E, 0x1E, 0x1E)
PURPLE_TIP   = RGBColor(0x86, 0x19, 0x8F)
CODE_GREEN   = RGBColor(0x10, 0xB9, 0x81)
CODE_GRAY    = RGBColor(0xD4, 0xD4, 0xD4)
WARNING_TEXT = RGBColor(0x85, 0x6D, 0x04)

# Background hex strings
BG_PURPLE      = "7C3AED"
BG_PURPLE_LIGHT = "EDE9FE"
BG_GREEN_BANNER = "10B981"
BG_GREEN_TIP    = "ECFDF5"
BG_BEST        = "D4EDDA"
BG_VIBE        = "FDF4FF"
BG_WARNING     = "FFF3CD"
BG_CODE        = "1E1E1E"
BG_TABLE_ALT   = "F8F9FA"
BG_TABLE_HDR   = "7C3AED"

BADGE_COLORS = {
    "purple": (BG_VIBE, PURPLE_TIP),
    "green": (BG_GREEN_TIP, GREEN_DARK),
    "blue": ("EBF5FB", RGBColor(0x1A, 0x5A, 0x96)),
}

# ─── Helpers ───

def set_paragraph_bg(paragraph, hex_color):
    """Set paragraph background/shading."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    paragraph._p.get_or_add_pPr().append(shd)

def set_cell_bg(cell, hex_color):
    """Set table cell background."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def add_run(paragraph, text, font_name=None, size=None, color=None, bold=None):
    """Add a styled run to a paragraph."""
    run = paragraph.add_run(text)
    if font_name:
        run.font.name = font_name
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    return run

def add_styled_paragraph(doc, text, size=12, color=None, bold=False, align=None, bg=None, space_after=None):
    """Add a paragraph with consistent styling."""
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if bg:
        set_paragraph_bg(p, bg)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text, size=size, color=color, bold=bold)
    return p

# ─── Document Building ───

def add_cover_page(doc, config):
    """Create the cover page."""
    # Blank line
    doc.add_paragraph()
    
    # Title
    add_styled_paragraph(doc, config["title"], size=40, color=PURPLE, bold=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Subtitle
    add_styled_paragraph(doc, config.get("subtitle", ""), size=26, color=GRAY_500,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Description
    add_styled_paragraph(doc, config.get("description", ""), size=18, color=GRAY_400,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Divider
    add_styled_paragraph(doc, "━" * 36, size=14, color=GRAY_200,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Tech stack
    if "tech_stack" in config:
        add_styled_paragraph(doc, "Technology Stack", size=14, color=GRAY_700, bold=True,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
        for line in config["tech_stack"]:
            add_styled_paragraph(doc, line, size=12, color=GRAY_500,
                                align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Duration/team badge
    if "duration" in config or "team_size" in config:
        parts = []
        if "duration" in config:
            parts.append(f"⏱️ Duration: {config['duration']}")
        if "team_size" in config:
            parts.append(f"👥 Team Size: {config['team_size']}")
        badge_text = "  " + "  •  ".join(parts) + "  "
        add_styled_paragraph(doc, badge_text, size=12, color=PURPLE_DARK,
                            align=WD_ALIGN_PARAGRAPH.CENTER, bg=BG_PURPLE_LIGHT)
    
    # Features line
    if "features" in config:
        add_styled_paragraph(doc, config["features"], size=11, color=GRAY_500,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Green banner
    if "banner" in config:
        add_styled_paragraph(doc, f"  {config['banner']}  ", size=14, color=WHITE, bold=True,
                            align=WD_ALIGN_PARAGRAPH.CENTER, bg=BG_GREEN_BANNER)
    
    doc.add_page_break()

def add_toc(doc, config):
    """Add table of contents."""
    add_styled_paragraph(doc, "📑 Table of Contents", size=20, color=GRAY_900, bold=True)
    
    page = 3  # Starting page estimate
    for part in config.get("parts", []):
        # Part entry (bold)
        p = doc.add_paragraph()
        add_run(p, part["title"], size=13, bold=True)
        add_run(p, f"\t{page}", size=13)
        page += 1
        
        for module in part.get("modules", []):
            p = doc.add_paragraph()
            add_run(p, module["title"], size=12)
            add_run(p, f"\t{page}", size=12)
            page += 3  # Estimate pages per module
        
        doc.add_paragraph()  # Spacing
    
    # Appendix
    if config.get("summary_cards"):
        p = doc.add_paragraph()
        add_run(p, "Appendix: Quick Reference", size=13, bold=True)
        add_run(p, f"\t{page}", size=13)
    
    doc.add_page_break()

def add_part_header(doc, title, subtitle=None):
    """Add a purple part header banner."""
    add_styled_paragraph(doc, f"  {title}  ", size=22, color=WHITE, bold=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER, bg=BG_PURPLE)
    if subtitle:
        add_styled_paragraph(doc, subtitle, size=15, color=GRAY_500,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

def add_module_header(doc, title, duration=None, fmt=None, badge_color="purple"):
    """Add module heading with duration badge."""
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.size = Pt(18)
        run.bold = True
    
    if duration or fmt:
        bg_hex, text_clr = BADGE_COLORS.get(badge_color, BADGE_COLORS["purple"])
        parts = []
        if duration:
            parts.append(f"⏱️ Duration: {duration}")
        if fmt:
            parts.append(f"🎯 Format: {fmt}")
        badge = " | ".join(parts)
        add_styled_paragraph(doc, badge, size=11, color=text_clr, bg=bg_hex)

def add_section_header(doc, title):
    """Add H2 section heading."""
    h = doc.add_heading(title, level=2)
    for run in h.runs:
        run.font.size = Pt(14)
        run.bold = True

def add_body_text(doc, text):
    """Add justified body text."""
    p = add_styled_paragraph(doc, text, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    return p

def add_bullets(doc, items):
    """Add bullet list."""
    for item in items:
        p = doc.add_paragraph(item, style='List Paragraph')
        for run in p.runs:
            run.font.size = Pt(12)

def add_vibe_box(doc, text):
    """Add purple vibe/tip callout box."""
    add_styled_paragraph(doc, text, size=12, color=PURPLE_TIP, bg=BG_VIBE)

def add_best_practice(doc, text):
    """Add green best practice box."""
    add_styled_paragraph(doc, text, size=12, color=GREEN_TEXT, bg=BG_BEST)

def add_warning_box(doc, text):
    """Add yellow warning box."""
    add_styled_paragraph(doc, text, size=12, color=WARNING_TEXT, bg=BG_WARNING)

def add_code_block(doc, code):
    """Add dark code block."""
    add_styled_paragraph(doc, code, size=11, color=CODE_GRAY, bg=BG_CODE)

def add_ai_prompt_box(doc, title, prompt, tip=None):
    """Add AI prompt box (green header + dark code + light tip)."""
    table = doc.add_table(rows=2 if not tip else 3, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row (green)
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, BG_GREEN_BANNER)
    p = cell.paragraphs[0]
    add_run(p, f"🤖 AI PROMPT: {title}", size=13, color=WHITE, bold=True)
    
    # Code row (dark)
    cell = table.rows[1].cells[0]
    set_cell_bg(cell, BG_CODE)
    p = cell.paragraphs[0]
    add_run(p, ">>> ", font_name="Consolas", size=12, color=CODE_GREEN, bold=True)
    add_run(p, prompt, font_name="Consolas", size=12, color=CODE_GRAY)
    
    # Tip row (light green)
    if tip:
        cell = table.rows[2].cells[0]
        set_cell_bg(cell, BG_GREEN_TIP)
        p = cell.paragraphs[0]
        add_run(p, "💬 ", size=11)
        add_run(p, tip, size=11, color=GREEN_DARK)
    
    doc.add_paragraph()  # Spacing

def add_quick_ref_box(doc, title, prompt):
    """Add compact quick reference prompt box (for appendix)."""
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, BG_GREEN_BANNER)
    p = cell.paragraphs[0]
    add_run(p, f"🤖 AI PROMPT: {title}", size=12, color=WHITE, bold=True)
    
    cell = table.rows[1].cells[0]
    set_cell_bg(cell, BG_CODE)
    p = cell.paragraphs[0]
    add_run(p, prompt, font_name="Consolas", size=11, color=CODE_GRAY)
    
    doc.add_paragraph()

def add_data_table(doc, headers, rows, numbered=False):
    """Add a styled data table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_bg(cell, BG_TABLE_HDR)
        p = cell.paragraphs[0]
        add_run(p, header, size=11, color=WHITE, bold=True)
    
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            if r_idx % 2 == 1:
                set_cell_bg(cell, BG_TABLE_ALT)
            p = cell.paragraphs[0]
            if numbered and c_idx == 0:
                add_run(p, str(val), size=11, color=PURPLE, bold=True)
            else:
                add_run(p, str(val), size=11)
    
    doc.add_paragraph()

def add_summary_cards(doc, cards):
    """Add side-by-side summary cards at the end."""
    if not cards or len(cards) < 2:
        return
    
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for i, card in enumerate(cards):
        cell = table.rows[0].cells[i]
        set_cell_bg(cell, BG_PURPLE_LIGHT)
        p = cell.paragraphs[0]
        add_run(p, card["title"] + "\n", size=13, color=PURPLE, bold=True)
        for item in card.get("items", []):
            add_run(p, f"• {item}\n", size=11, color=GRAY_700)

# ─── Content Renderer ───

def render_content(doc, content_items):
    """Render a list of content items."""
    for item in content_items:
        t = item["type"]
        if t == "text":
            add_body_text(doc, item["value"])
        elif t == "vibe":
            add_vibe_box(doc, item["value"])
        elif t == "best_practice":
            add_best_practice(doc, item["value"])
        elif t == "warning":
            add_warning_box(doc, item["value"])
        elif t == "bullets":
            add_bullets(doc, item["items"])
        elif t == "code":
            add_code_block(doc, item["value"])
        elif t == "ai_prompt":
            add_ai_prompt_box(doc, item["title"], item["prompt"], item.get("tip"))
        elif t == "quick_ref":
            add_quick_ref_box(doc, item["title"], item["prompt"])
        elif t == "table":
            add_data_table(doc, item["headers"], item["rows"])
        elif t == "numbered_table":
            add_data_table(doc, item["headers"], item["rows"], numbered=True)
        elif t == "page_break":
            doc.add_page_break()
        elif t == "heading":
            add_section_header(doc, item["value"])
        elif t == "subheading":
            add_styled_paragraph(doc, item["value"], size=12, bold=True)

# ─── Main ───

def generate_workshop(config, output_path):
    """Generate complete workshop document from config."""
    doc = Document()
    
    # Page setup (A4)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    
    # Cover page
    add_cover_page(doc, config)
    
    # Table of Contents
    add_toc(doc, config)
    
    # Parts & Modules
    for part in config.get("parts", []):
        add_part_header(doc, part["title"], part.get("subtitle"))
        
        for module in part.get("modules", []):
            add_module_header(doc, module["title"],
                            module.get("duration"),
                            module.get("format"),
                            module.get("badge_color", "purple"))
            
            for section in module.get("sections", []):
                add_section_header(doc, section["heading"])
                render_content(doc, section.get("content", []))
            
            doc.add_page_break()
    
    # Summary cards
    if config.get("summary_cards"):
        add_summary_cards(doc, config["summary_cards"])
    
    doc.save(output_path)
    print(f"✅ Generated: {output_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python3 workshop_docx_generator.py config.json output.docx")
        print("\nOr pass JSON via stdin:")
        print("  cat config.json | python3 workshop_docx_generator.py - output.docx")
        sys.exit(1)
    
    config_path = sys.argv[1]
    output_path = sys.argv[2]
    
    if config_path == "-":
        config = json.load(sys.stdin)
    else:
        with open(config_path) as f:
            config = json.load(f)
    
    generate_workshop(config, output_path)
