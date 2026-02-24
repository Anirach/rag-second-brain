#!/usr/bin/env python3
"""
Research Digest Generator - Creates weekly summary of relevant papers
Uses python-docx directly for professional formatting
"""

import json
import glob
import os
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def setup_styles(doc):
    """Setup professional document styles"""
    styles = doc.styles
    
    # Title Style
    if 'Digest Title' not in [s.name for s in styles]:
        title_style = styles.add_style('Digest Title', WD_STYLE_TYPE.PARAGRAPH)
        title_style.font.name = 'Arial'
        title_style.font.size = Pt(24)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0, 51, 102)
        title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_style.paragraph_format.space_after = Pt(6)
    
    # Subtitle Style
    if 'Digest Subtitle' not in [s.name for s in styles]:
        subtitle_style = styles.add_style('Digest Subtitle', WD_STYLE_TYPE.PARAGRAPH)
        subtitle_style.font.name = 'Arial'
        subtitle_style.font.size = Pt(14)
        subtitle_style.font.color.rgb = RGBColor(70, 70, 70)
        subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_style.paragraph_format.space_after = Pt(24)
    
    # Section Header
    if 'Section Header' not in [s.name for s in styles]:
        section_style = styles.add_style('Section Header', WD_STYLE_TYPE.PARAGRAPH)
        section_style.font.name = 'Arial'
        section_style.font.size = Pt(14)
        section_style.font.bold = True
        section_style.font.color.rgb = RGBColor(0, 51, 102)
        section_style.paragraph_format.space_before = Pt(18)
        section_style.paragraph_format.space_after = Pt(8)
    
    # Paper Title
    if 'Paper Title' not in [s.name for s in styles]:
        paper_style = styles.add_style('Paper Title', WD_STYLE_TYPE.PARAGRAPH)
        paper_style.font.name = 'Arial'
        paper_style.font.size = Pt(11)
        paper_style.font.bold = True
        paper_style.font.color.rgb = RGBColor(0, 51, 102)
        paper_style.paragraph_format.space_before = Pt(12)
        paper_style.paragraph_format.space_after = Pt(4)
    
    # Body Style
    if 'Digest Body' not in [s.name for s in styles]:
        body_style = styles.add_style('Digest Body', WD_STYLE_TYPE.PARAGRAPH)
        body_style.font.name = 'Arial'
        body_style.font.size = Pt(10)
        body_style.paragraph_format.line_spacing = 1.15
        body_style.paragraph_format.space_after = Pt(4)
        body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Metadata Style
    if 'Paper Metadata' not in [s.name for s in styles]:
        meta_style = styles.add_style('Paper Metadata', WD_STYLE_TYPE.PARAGRAPH)
        meta_style.font.name = 'Arial'
        meta_style.font.size = Pt(9)
        meta_style.font.italic = True
        meta_style.font.color.rgb = RGBColor(100, 100, 100)
        meta_style.paragraph_format.space_after = Pt(4)

def add_hyperlink(paragraph, text, url):
    """Add a hyperlink to a paragraph"""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    # Blue color
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0066CC')
    rPr.append(color)
    
    # Underline
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    
    paragraph._p.append(hyperlink)

def add_page_numbers(doc):
    """Add page numbers to footer"""
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # "Page X of Y" format
    run = para.add_run("Page ")
    run.font.size = Pt(9)
    
    # Current page
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText1 = OxmlElement('w:instrText')
    instrText1.text = "PAGE"
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText1)
    run._r.append(fldChar2)
    
    run2 = para.add_run(" of ")
    run2.font.size = Pt(9)
    
    # Total pages
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'begin')
    
    instrText2 = OxmlElement('w:instrText')
    instrText2.text = "NUMPAGES"
    
    fldChar4 = OxmlElement('w:fldChar')
    fldChar4.set(qn('w:fldCharType'), 'end')
    
    run2._r.append(fldChar3)
    run2._r.append(instrText2)
    run2._r.append(fldChar4)

def load_recent_papers(days_back=7):
    """Load papers from the last week"""
    all_papers = []
    cutoff_date = datetime.now() - timedelta(days=days_back)
    
    paper_files = glob.glob('/home/clawdbot/clawd/research/papers/papers_*.json')
    
    for file_path in paper_files:
        try:
            with open(file_path, 'r') as f:
                data = json.load(f)
                
            # Check if file is within date range
            timestamp = datetime.strptime(data['timestamp'], '%Y%m%d_%H%M%S')
            if timestamp >= cutoff_date:
                all_papers.extend(data['relevant_papers'])
                
        except Exception as e:
            print(f"Error processing {file_path}: {e}")
    
    # Deduplicate by arxiv_id
    seen = set()
    unique_papers = []
    for paper in all_papers:
        if paper['arxiv_id'] not in seen:
            seen.add(paper['arxiv_id'])
            unique_papers.append(paper)
    
    return unique_papers

def categorize_papers(papers):
    """Group papers by category"""
    categories = {}
    for paper in papers:
        cat = paper.get('category', 'Other')
        if cat not in categories:
            categories[cat] = []
        categories[cat].append(paper)
    return categories

def create_research_digest():
    """Generate comprehensive research digest"""
    papers = load_recent_papers(7)
    
    if not papers:
        print("📭 No relevant papers found in the last week")
        return None, 0, {}
    
    # Sort papers by relevance score
    papers.sort(key=lambda x: x.get('relevance_score', 0), reverse=True)
    
    # Take top papers for digest
    top_papers = papers[:15]
    
    # Categorize
    by_category = categorize_papers(top_papers)
    
    # Create document
    doc = Document()
    setup_styles(doc)
    add_page_numbers(doc)
    
    # Set margins
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    
    # Title
    week_start = (datetime.now() - timedelta(days=7)).strftime('%B %d')
    week_end = datetime.now().strftime('%B %d, %Y')
    
    title = doc.add_paragraph("Weekly Research Digest", style='Digest Title')
    subtitle = doc.add_paragraph(f"AI & Machine Learning Research Update\n{week_start} - {week_end}", style='Digest Subtitle')
    
    # Author info
    author_para = doc.add_paragraph()
    author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = author_para.add_run("Prepared for: Anirach Mingkhwan | FITM, KMUTNB")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph()  # Spacing
    
    # Executive Summary
    doc.add_paragraph("Executive Summary", style='Section Header')
    summary = doc.add_paragraph(style='Digest Body')
    summary.add_run(f"This weekly digest summarizes {len(top_papers)} highly relevant papers from arXiv across artificial intelligence, machine learning, and related fields. Papers are selected based on relevance to current research interests including AI/ML, longevity healthcare, biomedical AI, and educational technology.")
    
    # Statistics
    doc.add_paragraph("Overview Statistics", style='Section Header')
    stats = doc.add_paragraph(style='Digest Body')
    stats.add_run(f"• Total unique papers processed: {len(papers)}\n")
    stats.add_run(f"• Featured in this digest: {len(top_papers)}\n")
    stats.add_run(f"• Categories covered: {len(by_category)}\n")
    stats.add_run(f"• Period: {week_start} - {week_end}")
    
    # Category breakdown
    doc.add_paragraph("Papers by Category", style='Section Header')
    for cat, cat_papers in sorted(by_category.items(), key=lambda x: -len(x[1])):
        cat_para = doc.add_paragraph(style='Digest Body')
        cat_para.add_run(f"• {cat}: {len(cat_papers)} paper(s)")
    
    # Featured Papers
    doc.add_paragraph()
    doc.add_paragraph("Featured Papers", style='Section Header')
    
    for i, paper in enumerate(top_papers, 1):
        # Paper number and title
        paper_title = doc.add_paragraph(f"{i}. {paper['title']}", style='Paper Title')
        
        # Authors and metadata
        meta = doc.add_paragraph(style='Paper Metadata')
        authors = paper['authors']
        if len(authors) > 100:
            authors = authors[:100] + "..."
        meta.add_run(f"{authors}\n")
        meta.add_run(f"arXiv:{paper['arxiv_id']} | {paper.get('category', 'N/A')} | Published: {paper['published']}")
        
        # Summary
        summary_text = paper['summary']
        if len(summary_text) > 500:
            summary_text = summary_text[:500] + "..."
        summary_para = doc.add_paragraph(summary_text, style='Digest Body')
        
        # Relevance
        if paper.get('relevance_reasons'):
            relevance = doc.add_paragraph(style='Digest Body')
            run = relevance.add_run("Relevance: ")
            run.bold = True
            run.font.size = Pt(9)
            reasons = paper['relevance_reasons'][:3]  # Limit to 3 reasons
            relevance.add_run("; ".join(reasons)).font.size = Pt(9)
        
        # Link
        link_para = doc.add_paragraph(style='Digest Body')
        link_para.add_run("Link: ")
        add_hyperlink(link_para, paper['url'], paper['url'])
    
    # References
    doc.add_page_break()
    doc.add_paragraph("Complete Bibliography", style='Section Header')
    
    for i, paper in enumerate(top_papers, 1):
        ref = doc.add_paragraph(style='Digest Body')
        authors = paper['authors']
        if len(authors) > 80:
            authors = authors[:80] + " et al."
        year = paper['published'][:4]
        ref.add_run(f"[{i}] {authors} ({year}). {paper['title']}. arXiv:{paper['arxiv_id']}")
    
    # Save
    timestamp = datetime.now().strftime("%Y%m%d")
    filename = f"/home/clawdbot/clawd/research/digests/research_digest_{timestamp}.docx"
    
    doc.save(filename)
    print(f"📊 Research digest created: {filename}")
    
    return filename, len(top_papers), by_category

def main():
    """Generate weekly research digest"""
    print("📚 Generating weekly research digest...")
    digest_file, paper_count, categories = create_research_digest()
    
    if digest_file:
        print(f"✅ Weekly digest generated: {digest_file}")
        print(f"📄 Papers included: {paper_count}")
        print(f"📁 Categories: {list(categories.keys())}")
        return digest_file, paper_count, categories
    else:
        print("❌ No digest generated - insufficient papers")
        return None, 0, {}

if __name__ == "__main__":
    main()
